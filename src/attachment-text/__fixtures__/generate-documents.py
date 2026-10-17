# Regenerates the PDF and DOCX fixtures used by the attachment-text tests.
# Run from the project root:
#   uv run --with pypdf --with python-docx --with cryptography python src/attachment-text/__fixtures__/generate-documents.py
#
# report.pdf: two pages of text. protected.pdf: the same document, password "secret".
# letter.docx: a heading, a paragraph and a two-row table.

from pathlib import Path

from docx import Document
from pypdf import PdfReader, PdfWriter

HERE = Path(__file__).parent


def pdf_with_pages(pages: list[str]) -> bytes:
    """A minimal valid PDF, one Helvetica text line per page, with a correct xref."""
    objects: list[bytes] = []
    page_ids = [3 + 2 * i for i in range(len(pages))]
    font_id = 3 + 2 * len(pages)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    for index, text in enumerate(pages):
        content_id = page_ids[index] + 1
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        stream = f"BT /F1 18 Tf 72 760 Td ({text}) Tj ET".encode()
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


report = pdf_with_pages(["Quarterly report 2026", "Revenue rose by 12 percent"])
(HERE / "report.pdf").write_bytes(report)

writer = PdfWriter(clone_from=PdfReader(HERE / "report.pdf"))
writer.encrypt(user_password="secret", owner_password="secret-owner", algorithm="AES-128")
with (HERE / "protected.pdf").open("wb") as handle:
    writer.write(handle)

letter = Document()
letter.add_heading("Bursary application", level=1)
letter.add_paragraph("Thank you for your application. Please find the schedule below.")
table = letter.add_table(rows=2, cols=2)
table.cell(0, 0).text = "Deadline"
table.cell(0, 1).text = "30 November 2026"
table.cell(1, 0).text = "Interview"
table.cell(1, 1).text = "7 December 2026"
letter.save(HERE / "letter.docx")

for name in ("report.pdf", "protected.pdf", "letter.docx"):
    print(f"wrote {name} ({(HERE / name).stat().st_size} bytes)")
